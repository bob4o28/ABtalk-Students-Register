import pandas as pd
from docx import Document
import os

# Load the Excel file
excel_path = "C:/Users/Boris/boris/Apps and softwares/ABtalk/Resurses/Students_2025_Register/DBABtalk_2025_АГ.xlsx"
df = pd.read_excel(excel_path)

# Load the Word template
template_path = "C:/Users/Boris/boris/Apps and softwares/ABtalk/Resurses/извинителна бележка до училища 2025.docx"

# Define the output folder
output_folder = "C:/Users/Boris/boris/Apps and softwares/ABtalk/Resurses/Students_2025_Register/WordFIles"
os.makedirs(output_folder, exist_ok=True)

# Loop through each student in the table
for index, row in df.iterrows():
    first_name = str(row[0]).strip()  # First Name (Column 1)
    last_name = str(row[1]).strip()   # Last Name (Column 2)
    class_number = str(row[4]).strip()  # Class Number (Column 5)
    class_letter = str(row[5]).strip()  # Class Letter (Column 6)
    
    full_name = f"{first_name} {last_name}"
    student_class = f"{class_number}{class_letter}"  # Example: "10A"

    # Open the Word document
    doc = Document(template_path)

    # Replace placeholders in the document
    for para in doc.paragraphs:
        if "NAMES" in para.text:
            para.text = para.text.replace("NAMES", f" {full_name}")
        if "от  клас" in para.text:
            para.text = para.text.replace("от  клас", f"от {student_class} клас")

    # Save the new document with the student's name
    output_path = os.path.join(output_folder, f"Excuse_{full_name}.docx")
    doc.save(output_path)

print("All files have been generated successfully!")
